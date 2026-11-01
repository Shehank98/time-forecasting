import io

import openpyxl
import pytest
from pptx import Presentation

import synth
from app.services import analytics_kpi, analytics_lmrb as A, analytics_tv as T, export_excel as X, export_pdf as P, export_pptx as PP
from app.services import forecast as F, ingest, ingest_kpi, narrative, report_data as R
from helpers import api, series_payload, summary, wait_job

PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99=\x1d\x00\x00\x00\x00IEND\xaeB`\x82", "image/png")
BRAND = {"agency_name": "Media Mind LK", "primary_color": "#112233", "report_footer": "Confidential", "disclaimer": "Forecasts are estimates."}


@pytest.fixture(scope="module")
def kpi_rep():
    p = ingest_kpi.propose(synth.to_csv_bytes(synth.make_kpi()), "k.csv")
    df, _, meta = ingest_kpi.apply(p["token"], p["suggested"])
    dash = analytics_kpi.dashboard(df, meta, df)
    fc = F.run(df, "kpi", {"metric": "leads", "entity_type": "segment", "granularity": "weekly", "horizon_days": 56}, meta=meta)
    sc = F.run_scenarios(df, "kpi", {"target": "leads", "driver": "media_spend_lkr", "granularity": "weekly", "horizon_days": 28, "multipliers": [0.8, 1, 1.2], "ratio_label": "CPL"}, meta)
    pc = F.pacing(df, "kpi", {"metric": "leads", "entity_type": "total", "start": "2026-09-01", "end": "2026-10-15", "target": 5000}, meta)
    return R.build("kpi", dash, {}, ["k.csv"], fc, None, ["Segment: All"], client={"name": "Acme Foods", "brand_color": "#0E9F6E"}, branding=BRAND,
                   scenarios=sc, pacing=pc, dataset_meta=meta, dataset_name="Leads", title="Q4 planning report", logos={"agency": PNG, "client": PNG})


def test_rule_based_story_is_specific(kpi_rep):
    st = narrative.story(kpi_rep)
    assert st["source"] == "rules" and st["headline"]
    text = " ".join(st["findings"] + st["actions"])
    assert "Current pace" in text and "CPL" in text and "Pacing" in text
    assert all("—" not in t for t in st["findings"] + st["actions"] + st["watch"])


def test_story_prefers_gemini_when_present(kpi_rep):
    import copy
    rep = copy.copy(kpi_rep)
    rep.ai = {"executive_summary": {"statement": "AI headline", "confidence": "high", "evidence": []}, "recommended_actions": [{"statement": "Do X", "confidence": "low", "evidence": []}],
              "competitor_moves": [], "predicted_behaviour": [], "threats": [], "opportunities": [], "data_caveats": ["c"]}
    st = narrative.story(rep)
    assert st["source"] == "gemini" and st["headline"] == "AI headline" and st["actions"] == ["Do X"]


def test_pdf_branded_and_complete(kpi_rep):
    pdf = P.build_pdf(kpi_rep)
    assert pdf[:5] == b"%PDF-" and pdf.count(b"/Type /Page") >= 5


def test_pptx_has_native_charts_and_sections(kpi_rep):
    prs = Presentation(io.BytesIO(PP.build_pptx(kpi_rep)))
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    assert len(prs.slides) >= 9 and charts >= 5
    assert any("Executive summary" in t for t in texts) and any("Budget scenarios" in t for t in texts) and any("Pacing" in t for t in texts)
    assert "Q4 planning report" in texts[0] and "Acme Foods" in texts[0]


def test_excel_scenarios_and_pacing_formulas(kpi_rep, tmp_path):
    wb = X.build_workbook(kpi_rep)
    assert {"Cover", "KPIs", "Scenarios", "Pacing tracker", "Methodology"} <= set(wb.sheetnames)
    sc = wb["Scenarios"]
    assert str(sc["F5"].value).startswith("=IF(D5>0,B5/D5") and any(str(c.value).startswith("=FORECAST(") for row in sc.iter_rows() for c in row)
    pt = wb["Pacing tracker"]
    assert pt["B3"].value == 5000 and str(pt["B9"].value).startswith("=$B$3/") and str(pt["G9"].value).startswith("=SUM(")
    assert wb["Cover"]["A4"].value == "Q4 planning report" and wb["Cover"]._images
    assert wb["KPIs"]["A3"].fill.fgColor.rgb.endswith("112233")
    p = tmp_path / "r.xlsx"; wb.save(p)
    assert openpyxl.load_workbook(p)["Scenarios"]["A5"].value == "Budget -20%"


@pytest.mark.parametrize("kind", ["tv", "lmrb"])
def test_all_formats_for_tv_and_lmrb(kind):
    lm = ingest.parse(synth.to_csv_bytes(synth.make_lmrb()), "l.csv", "lmrb").df
    tv = ingest.parse(synth.to_csv_bytes(synth.make_tv()), "t.csv", "tv").df
    dash = A.dashboard(lm, None, None, "Nestle Lanka Ltd") if kind == "lmrb" else T.dashboard(tv, lm)
    fc = F.run(lm if kind == "lmrb" else tv, kind, {"metric": "spend" if kind == "lmrb" else "tvr", "entity_type": "advertiser" if kind == "lmrb" else "timebelt", "granularity": "weekly", "horizon_days": 28})
    rep = R.build(kind, dash, {}, ["x.csv"], fc, None, [], client={"name": "Acme"}, branding=BRAND)
    assert P.build_pdf(rep)[:5] == b"%PDF-"
    assert len(Presentation(io.BytesIO(PP.build_pptx(rep))).slides) >= 5
    X.build_workbook(rep)


def test_export_endpoint_all_formats_saved_per_client():
    c = api()
    dash, meta = summary("lmrb")
    fc = wait_job(c, c.post("/api/forecast/run", json={"series": series_payload(gran="weekly"), "params": {"granularity": "weekly", "horizon_days": 28}}).json()["job_id"])["result"]
    daily = series_payload(metric="spend", entity_type="total")[0]
    pace = c.post("/api/pacing/run", json={"series": daily, "params": {"start": "2026-09-01", "end": "2026-09-30", "target": 3e7}}).json()
    body = {"kind": "lmrb", "dash": dash, "meta": meta, "forecast": fc, "pacing": pace, "evidence": ["Unilever burst expected next week."]}
    for fmt, magic in (("pdf", b"%PDF-"), ("pptx", b"PK"), ("xlsx", b"PK")):
        r = c.post("/api/export", json={**body, "format": fmt, "title": "Monthly competitor review"})
        assert r.status_code == 200 and r.content[:len(magic)] == magic and r.headers["x-report-id"]
    reps = c.get("/api/reports").json()
    assert [x["format"] for x in reps] == ["xlsx", "pptx", "pdf"]
    dl = c.get(f"/api/reports/{reps[0]['id']}/download")
    assert dl.status_code == 200 and dl.content[:2] == b"PK"
    other = api()
    assert other.get(f"/api/reports/{reps[0]['id']}/download").status_code == 404     # another client cannot fetch it


@pytest.mark.parametrize("kind", ["tv", "kpi"])
def test_export_endpoint_tv_and_kpi(kind):
    c = api()
    dash, meta = summary(kind)
    r = c.post("/api/export", json={"kind": kind, "format": "pptx", "dash": dash, "meta": meta, "dataset": "Leads" if kind == "kpi" else ""})
    assert r.status_code == 200 and r.content[:2] == b"PK"
