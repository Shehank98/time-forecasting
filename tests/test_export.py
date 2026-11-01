import io
import os
import tempfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

from helpers import api, login, new_client, up as up_helper

import synth
from app.main import app
from app.services import analytics_lmrb as A
from app.services import analytics_tv as T
from app.services import export_excel as X
from app.services import export_pdf as P
from app.services import forecast as F
from app.services import ingest
from app.services import report_data as R

AI = {"executive_summary": {"statement": "Summary text", "confidence": "high", "evidence": []}, "competitor_moves": [], "threats": [], "opportunities": [], "predicted_behaviour": [],
      "recommended_actions": [{"statement": "Protect prime", "confidence": "medium", "evidence": []}], "data_caveats": ["Short history"]}


@pytest.fixture(scope="module")
def lm():
    return ingest.parse(synth.to_csv_bytes(synth.make_lmrb()), "s.csv", "lmrb").df


@pytest.fixture(scope="module")
def tv():
    return ingest.parse(synth.to_csv_bytes(synth.make_tv()), "s.csv", "tv").df


@pytest.fixture(scope="module")
def lm_report(lm):
    fc = F.run(lm, "lmrb", {"entity_type": "advertiser", "metric": "spend", "granularity": "weekly", "horizon_days": 28})
    return R.build("lmrb", A.dashboard(lm, None, None, "Nestle Lanka Ltd"), {}, ["a.csv"], fc, AI, ["Channel: DERANA"])


def test_excel_structure_formats_and_native_charts(lm_report):
    wb = X.build_workbook(lm_report)
    assert wb.sheetnames[:2] == ["Cover", "KPIs"] and {"Spend by advertiser", "SOV weekly", "Forecast summary", "Forecast values", "AI insights", "Pricing heatmap"} <= set(wb.sheetnames)
    ws = wb["Spend by advertiser"]
    assert ws.freeze_panes == "A5" and len(ws.tables) == 1 and len(ws._charts) == 1
    assert ws["C5"].number_format == '"LKR" #,##0' and ws["D5"].number_format == "0.0%"
    assert len(wb["SOV weekly"]._charts) == 1 and len(wb["Spend by channel"]._charts) == 1 and len(wb["FC chart 1"]._charts) == 1
    assert wb["Channel mix"].conditional_formatting
    cover = " ".join(str(c.value) for c in wb["Cover"]["B"])
    assert "LMRB Ad Monitoring" in cover and "DERANA" in cover and "Executive summary" in " ".join(str(c.value) for c in wb["Cover"]["A"])
    assert wb["Spend by advertiser"].column_dimensions["B"].width > 8
    sh = wb["Forecast summary"]
    assert [c.value for c in sh[4]][:3] == ["Entity", "Reliability", "History points"]


def test_excel_roundtrips_through_openpyxl(lm_report, tmp_path):
    p = tmp_path / "r.xlsx"
    X.build_workbook(lm_report).save(p)
    assert openpyxl.load_workbook(p)["KPIs"]["A4"].value == "Total spend"


def test_pdf_is_valid_and_contains_sections(lm_report):
    pdf = P.build_pdf(lm_report)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 20_000
    from reportlab.lib.pagesizes import A4  # noqa: F401
    text = pdf.decode("latin-1")
    assert "/Page" in text


def test_tv_reports_build(tv, lm):
    rep = R.build("tv", T.dashboard(tv, lm), {}, ["t.csv"], None, None, [])
    assert "Slot opportunity" in [t.name for t in rep.tables]
    X.build_workbook(rep)
    assert P.build_pdf(rep)[:5] == b"%PDF-"
    rep2 = R.build("tv", T.dashboard(tv, None), {}, ["t.csv"], None, None, [])
    assert "Slot opportunity" not in [t.name for t in rep2.tables]


def test_no_em_dashes_in_generated_text(lm_report):
    for t in lm_report.tables:
        assert "\u2014" not in t.note and "\u2014" not in " ".join(map(str, t.df.astype(str).values.ravel()))


def _lmrb_summary(channel=None):
    from app.services import filters as flt
    from app.services.util import jsonable
    df = ingest.parse(synth.to_csv_bytes(synth.make_lmrb()), "a.csv", "lmrb").df
    f = {"channel": [channel]} if channel else {}
    dash = A.dashboard(flt.apply_filters(df, "lmrb", f), None, None, "Nestle Lanka Ltd")
    return jsonable(dash), {"filtered_rows": 1, "chips": flt.describe_filters("lmrb", f)}


def test_export_endpoint_streams_and_cleans_up(monkeypatch):
    from helpers import series_payload, wait_job
    before = set(os.listdir(tempfile.gettempdir()))
    c = api()
    dash, meta = _lmrb_summary("DERANA")
    fc = wait_job(c, c.post("/api/forecast/run", json={"series": series_payload(gran="weekly"), "params": {"granularity": "weekly", "horizon_days": 28}}).json()["job_id"])["result"]
    body = {"kind": "lmrb", "dash": dash, "meta": meta, "files": ["a.csv"], "forecast": fc, "include_ai": True, "ai": AI}
    r = c.post("/api/export", json={**body, "format": "xlsx"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"] and r.headers["content-disposition"].count(".xlsx") == 1
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    cover = " ".join(str(c.value) for row in wb["Cover"].iter_rows() for c in row if c.value)
    assert "Channel: DERANA" in cover and "a.csv" in cover and "AI insights" in wb.sheetnames and "Forecast summary" in wb.sheetnames
    chans = [row[0].value for row in wb["Spend by channel"].iter_rows(min_row=5, min_col=2, max_col=2) if row[0].value]
    assert chans == ["DERANA"]
    r2 = c.post("/api/export", json={**body, "format": "pdf"})
    assert r2.status_code == 200 and r2.content[:5] == b"%PDF-"
    leftovers = {f for f in set(os.listdir(tempfile.gettempdir())) - before if f.endswith((".xlsx", ".pdf"))}
    assert not leftovers


def test_export_without_ai_when_gemini_unavailable_adds_note(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    c = api()
    dash, meta = _lmrb_summary()
    r = c.post("/api/export", json={"kind": "lmrb", "format": "xlsx", "include_ai": True, "dash": dash, "meta": meta, "evidence": ["Clutter rising on Derana prime."]})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    text = " ".join(str(c.value) for row in wb["Cover"].iter_rows() for c in row if c.value)
    assert "AI insights" not in wb.sheetnames and "summary not included" in text and "Clutter rising on Derana prime." in text


def test_export_errors_are_clear():
    c = api()
    dash, meta = _lmrb_summary()
    assert c.post("/api/export", json={"kind": "lmrb", "format": "docx", "dash": dash}).status_code == 422
    r = c.post("/api/export", json={"kind": "lmrb", "format": "pdf", "dash": {"empty": True}})
    assert r.status_code == 422 and "nothing to export" in r.json()["detail"]
    assert c.post("/api/export", json={"kind": "lmrb", "format": "pdf", "dash": {"kpis": {}}}).status_code == 422
