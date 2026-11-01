"""Client deck (.pptx) with native, editable charts: one message per slide."""
from __future__ import annotations

import io
import logging
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .narrative import METHODOLOGY, fmt_value, story
from .report_data import Report

log = logging.getLogger(__name__)

W, H = Inches(13.333), Inches(7.5)
INK, MUTED, LIGHT = RGBColor(0x1B, 0x1D, 0x36), RGBColor(0x64, 0x6A, 0x8A), RGBColor(0xF4, 0xF5, 0xFB)
PALETTE = ["5B5CE6", "12A6A0", "F08A24", "D6457F", "3B9AE1", "8B5CF6", "C9A30A", "3EA65A", "E5484D", "0B8FB8"]


def _rgb(hex_: str) -> RGBColor:
    h = (hex_ or "#4F46E5").lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class Deck:
    def __init__(self, rep: Report) -> None:
        self.rep = rep
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.brand = _rgb(rep.branding.get("primary_color"))
        self.accent = _rgb(rep.client.get("brand_color"))
        self.agency = rep.branding.get("agency_name") or "Media Radar"
        self.client = rep.client.get("name") or "Client"
        self.n = 0

    # ---------------- primitives
    def slide(self, title: str, subtitle: str = "") -> Any:
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = self.accent; bar.line.fill.background()
        self.text(s, title, Inches(0.6), Inches(0.35), Inches(11), Inches(0.7), 26, True, self.brand)
        if subtitle:
            self.text(s, subtitle, Inches(0.6), Inches(1.0), Inches(12), Inches(0.5), 12, False, MUTED)
        logo = self.rep.logos.get("client")
        if logo:
            self.picture(s, logo, Inches(11.9), Inches(0.3), Inches(0.6))
        self.text(s, f"{self.agency}  |  {self.client}  |  {self.rep.branding.get('report_footer', '')}", Inches(0.6), Inches(7.05), Inches(10.5), Inches(0.3), 9, False, MUTED)
        self.text(s, str(self.n), Inches(12.3), Inches(7.05), Inches(0.5), Inches(0.3), 9, False, MUTED, PP_ALIGN.RIGHT)
        return s

    def text(self, s, txt: str, x, y, w, h, size: int = 14, bold: bool = False, color: RGBColor = INK, align=PP_ALIGN.LEFT):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = align
        r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        return tb

    def bullets(self, s, items: list[str], x, y, w, h, size: int = 15) -> None:
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame; tf.word_wrap = True
        for i, t in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            r = p.add_run(); r.text = "•  " + t; r.font.size = Pt(size); r.font.color.rgb = INK

    def picture(self, s, logo: tuple[bytes, str], x, y, h) -> None:
        try:
            s.shapes.add_picture(io.BytesIO(logo[0]), x, y, height=h)
        except Exception:  # noqa: BLE001 - a bad logo must not break the deck
            log.warning("Logo skipped")

    def tile(self, s, label: str, value: str, x, y, w=Inches(3.0), h=Inches(1.1)) -> None:
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h); box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = RGBColor(0xE5, 0xE7, 0xF2)
        self.text(s, label.upper(), x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), Inches(0.3), 10, False, MUTED)
        self.text(s, value, x + Inches(0.15), y + Inches(0.42), w - Inches(0.3), Inches(0.6), 20, True, INK)

    def table(self, s, rows: list[list[str]], x, y, w, row_h=Inches(0.36), size: int = 11) -> None:
        shp = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
        t = shp.table
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = t.cell(i, j); c.text = str(v)
                para = c.text_frame.paragraphs[0]; para.font.size = Pt(size); para.font.bold = i == 0
                if i == 0:
                    c.fill.solid(); c.fill.fore_color.rgb = self.brand; para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    def chart(self, s, kind, cats: list[str], series: list[tuple[str, list]], x, y, w, h, legend: bool = True, number_format: str = "#,##0"):
        cd = CategoryChartData(); cd.categories = cats
        for name, vals in series:
            cd.add_series(name, [None if v is None else float(v) for v in vals])
        gf = s.shapes.add_chart(kind, x, y, w, h, cd)
        ch = gf.chart
        ch.has_legend = legend and len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(10)
        try:
            ch.value_axis.tick_labels.font.size = Pt(10); ch.value_axis.tick_labels.number_format = number_format; ch.value_axis.tick_labels.number_format_is_linked = False
            ch.category_axis.tick_labels.font.size = Pt(9); ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xE5, 0xE7, 0xF2)
        except (AttributeError, ValueError):
            pass
        for i, ser in enumerate(ch.plots[0].series):
            fill = _rgb(PALETTE[i % len(PALETTE)])
            if kind in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
                ser.format.line.color.rgb = fill; ser.format.line.width = Pt(2.25); ser.smooth = False
            elif kind != XL_CHART_TYPE.PIE:
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = fill
        return ch

    def forecast_chart(self, s, r: dict, x, y, w, h) -> None:
        hd, fd = r["history"]["dates"][-40:], r["forecast"]["dates"]
        hv = list(r["history"]["values"][-40:])
        n = len(hd)
        pad = lambda arr: [None] * (n - 1) + [hv[-1]] + list(arr)
        ch = self.chart(s, XL_CHART_TYPE.LINE, hd + fd, [("Actual", hv + [None] * len(fd)), ("P10", pad(r["forecast"]["p10"])), ("Forecast P50", pad(r["forecast"]["p50"])), ("P90", pad(r["forecast"]["p90"]))], x, y, w, h)
        sers = ch.plots[0].series
        style = [("1B1D36", 2.5, None), ("A5A6F6", 1.25, MSO_LINE_DASH_STYLE.DASH), ("5B5CE6", 2.5, MSO_LINE_DASH_STYLE.DASH), ("A5A6F6", 1.25, MSO_LINE_DASH_STYLE.DASH)]
        for ser, (col, wd, dash) in zip(sers, style):
            ser.format.line.color.rgb = _rgb(col); ser.format.line.width = Pt(wd)
            if dash:
                ser.format.line.dash_style = dash

    # ---------------- slides
    def build(self) -> bytes:
        rep, st = self.rep, story(self.rep)
        # title
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6]); self.n += 1
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H); bg.fill.solid(); bg.fill.fore_color.rgb = self.accent; bg.line.fill.background()
        if rep.logos.get("agency"):
            self.picture(s, rep.logos["agency"], Inches(0.9), Inches(0.6), Inches(0.8))
        else:
            self.text(s, self.agency, Inches(0.9), Inches(0.6), Inches(6), Inches(0.6), 18, True, self.brand)
        if rep.logos.get("client"):
            self.picture(s, rep.logos["client"], Inches(11.2), Inches(0.6), Inches(0.9))
        self.text(s, rep.title, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.2), 40, True, INK)
        self.text(s, f"Prepared for {self.client} by {self.agency}", Inches(0.9), Inches(3.8), Inches(11), Inches(0.6), 20, False, self.brand)
        self.text(s, f"{rep.card}   |   {rep.date_range}   |   {rep.generated}", Inches(0.9), Inches(4.5), Inches(11.5), Inches(0.5), 13, False, MUTED)
        # executive summary
        s = self.slide("Executive summary", "Written by Gemini from the numbers in this deck" if st["source"] == "gemini" else "Automated summary built from the numbers in this deck")
        self.text(s, st["headline"], Inches(0.6), Inches(1.5), Inches(12.1), Inches(1.0), 20, True, INK)
        for i, (lab, v, f) in enumerate(rep.kpis[:4]):
            self.tile(s, lab, _fmt(v, f), Inches(0.6 + i * 3.1), Inches(2.7))
        self.bullets(s, st["findings"][:4], Inches(0.6), Inches(4.0), Inches(12.1), Inches(2.9), 14)
        s = self.slide("Recommended actions", "What we suggest for the next flight")
        self.bullets(s, st["actions"][:6], Inches(0.6), Inches(1.6), Inches(12.1), Inches(3.6), 16)
        if st["watch"]:
            self.text(s, "Watch-outs", Inches(0.6), Inches(5.1), Inches(4), Inches(0.4), 14, True, self.brand)
            self.bullets(s, st["watch"][:3], Inches(0.6), Inches(5.5), Inches(12.1), Inches(1.4), 12)
        if rep.forecast:
            self.forecast_slides(rep.forecast)
        if rep.scenarios:
            self.scenario_slide(rep.scenarios)
        if rep.pacing:
            self.pacing_slide(rep.pacing)
        for sv in rep.saved:
            res = sv["result"]
            if sv["ftype"] == "forecast" and res.get("results"):
                s = self.slide(f"Saved forecast: {sv['title']}", f"Saved {str(sv['created_at'])[:10]}")
                self.forecast_chart(s, res["results"][0], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2))
            elif sv["ftype"] == "scenarios" and res.get("scenarios"):
                self.scenario_slide(res, f"Saved scenario plan: {sv['title']}")
            elif sv["ftype"] == "pacing" and res.get("cumulative"):
                self.pacing_slide(res, f"Saved pacing check: {sv['title']}")
        self.workspace_slides()
        # appendix
        s = self.slide("Appendix: methodology", "How the forecasts were made and how far to trust them")
        self.bullets(s, METHODOLOGY[:6], Inches(0.6), Inches(1.5), Inches(12.1), Inches(4.5), 12)
        self.text(s, rep.branding.get("disclaimer", ""), Inches(0.6), Inches(6.3), Inches(12.1), Inches(0.6), 10, False, MUTED)
        if rep.forecast:
            s = self.slide("Appendix: forecast reliability", "Backtest on the most recent periods against a seasonal naive baseline")
            rows = [["Entity", "Reliability", "History", "sMAPE", "MASE", "Inside low-high range", "Baseline sMAPE"]]
            for r in rep.forecast["results"][:12]:
                bt = r.get("backtest") or {}
                m, b = bt.get("model", {}), bt.get("baseline", {})
                rows.append([r["entity"], r["reliability"]["level"], f"{r['n_points']} {r['granularity']}", _p(m.get("smape")), "n/a" if m.get("mase") is None else f"{m['mase']:.2f}",
                             _p(m.get("coverage")), _p(b.get("smape"))])
            self.table(s, rows, Inches(0.6), Inches(1.6), Inches(12.1), size=11)
        buf = io.BytesIO(); self.prs.save(buf)
        return buf.getvalue()

    def forecast_slides(self, fc: dict) -> None:
        rs, p = fc["results"], fc["params"]
        label, unit = rs[0].get("metric_label", p.get("metric")), rs[0].get("unit", "number")
        if len(rs) > 1:
            s = self.slide(f"{label}: outlook", f"{len(rs)} {p.get('entity_type')}s, next 4 weeks, Forecast engine")
            rows = [["Entity", "Trend", "Change", "Expected", "Next peak", "Reliability"]]
            for r in sorted(rs, key=lambda r: -(r["outlook"]["trend_change_pct"] or 0))[:10]:
                o = r["outlook"]; pk = (r.get("peaks") or {}).get("peaks") or []
                rows.append([r["entity"], o["trend"].title(), _p(o.get("trend_change_pct"), True), fmt_value(o["expected_next_4w"] if o.get("additive") else o.get("avg_next_4w"), unit),
                             pk[0]["date"] if pk else "", r["reliability"]["level"]])
            self.table(s, rows, Inches(0.6), Inches(1.6), Inches(12.1), size=12)
        for r in rs[:4]:
            o = r["outlook"]
            pk = (r.get("peaks") or {}).get("peaks") or []
            s = self.slide(f"{r['entity']}: {label} forecast", f"{o['trend'].title()} ({_p(o.get('trend_change_pct'), True)} vs recent), reliability {r['reliability']['level']}. Dashed lines: P10, P50 and P90.")
            self.forecast_chart(s, r, Inches(0.6), Inches(1.6), Inches(8.6), Inches(5.2))
            notes = []
            exp = o["expected_next_4w"] if o.get("additive") else o.get("avg_next_4w")
            notes.append(f"Expected next {o['n_weeks']} {o.get('block', 'week')}s: {fmt_value(exp, unit)}")
            if pk:
                notes.append(f"Peak around {pk[0]['date']}. Be live by {pk[0]['launch_by']}.")
            notes += [b["text"] for b in (r.get("band_breaches") or [])[:2]]
            notes += [f"Bursts expected: {', '.join(o['burst_windows'])}"] if o.get("burst_windows") else []
            self.bullets(s, notes, Inches(9.4), Inches(1.8), Inches(3.4), Inches(5), 12)

    def scenario_slide(self, sc: dict, title: str = "Budget scenarios") -> None:
        rows_ = sc["scenarios"]
        ratio = sc["params"].get("ratio_label", "Cost per result")
        s = self.slide(title, f"Next {sc['params'].get('horizon_points')} {sc['params'].get('granularity')} periods. {'Forecast engine with spend as a covariate.' if sc.get('method') == 'timesfm_xreg' else 'Linear driver model with seasonal smoothing.'}")
        self.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, [r["scenario"] for r in rows_], [("Expected results (P50)", [r["kpi_p50"] for r in rows_])], Inches(0.6), Inches(1.6), Inches(5.6), Inches(5.0), legend=False)
        rows = [["Scenario", "Spend", "Results P50", "Range", ratio]]
        for r in rows_:
            rows.append([r["scenario"], f"LKR {r['driver_total'] / 1e6:,.1f}M", f"{r['kpi_p50']:,.0f}", f"{r['kpi_p10']:,.0f} to {r['kpi_p90']:,.0f}", f"LKR {r['ratio_p50']:,.0f}" if r.get("ratio_p50") else "n/a"])
        self.table(s, rows, Inches(6.5), Inches(1.7), Inches(6.3), size=11)

    def pacing_slide(self, pc: dict, title: str = "Pacing") -> None:
        s = self.slide(title, f"{pc['metric_label']} for {pc['entity']}, flight {pc['start']} to {pc['end']}: {pc['status']}")
        cum, pr = pc["cumulative"], pc["projection"]
        n = len(cum["dates"])
        lastv = cum["values"][-1] if cum["values"] else 0
        pad = lambda arr: [None] * (n - 1) + [lastv] + list(arr)
        self.chart(s, XL_CHART_TYPE.LINE, cum["dates"] + pr["dates"], [("Actual", list(cum["values"]) + [None] * len(pr["dates"])), ("Projection P50", pad(pr["p50"])),
                                                                      ("Target", [pc["target"]] * (n + len(pr["dates"])))], Inches(0.6), Inches(1.6), Inches(8.6), Inches(5.2))
        self.tile(s, "Delivered to date", f"{pc['actual_to_date']:,.0f}", Inches(9.5), Inches(1.7), Inches(3.3))
        self.tile(s, "Projected end (P50)", f"{pc['projected_p50']:,.0f}", Inches(9.5), Inches(3.0), Inches(3.3))
        self.tile(s, "Target", f"{pc['target']:,.0f}", Inches(9.5), Inches(4.3), Inches(3.3))
        if pc.get("required_per_day"):
            self.text(s, f"Needed per day: {pc['required_per_day']:,.0f} (recent {pc['recent_run_rate']:,.0f})", Inches(9.5), Inches(5.6), Inches(3.3), Inches(0.8), 12, False, MUTED)

    def workspace_slides(self) -> None:
        rep, d = self.rep, self.rep.dash
        if d.get("empty"):
            return
        if rep.kind == "lmrb":
            sw = d["sov_weekly"]
            if sw["weeks"]:
                s = self.slide("Share of voice over time", "Weekly share of spend by advertiser")
                self.chart(s, XL_CHART_TYPE.AREA_STACKED_100, sw["weeks"], [(a["name"], a["spend"]) for a in sw["advertisers"]], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), number_format="0%")
            top = d["by_advertiser"][:10]
            s = self.slide("Spend by advertiser", f"{d['window']['start']} to {d['window']['end']}")
            self.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [r["name"] for r in top][::-1], [("Spend (LKR)", [r["spend"] for r in top][::-1])], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), legend=False)
            g = d.get("grp")
            if g:
                s = self.slide("GRPs and cost per rating point", f"{g['total_grps']:,.1f} GRPs at an average CPRP of LKR {g['cprp']:,.0f}. {g['match_rate_pct']:.0f}% of spots matched to programme ratings.")
                ga = g["by_advertiser"][:8]
                self.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [r["name"] for r in ga][::-1], [("GRPs", [r["grps"] for r in ga][::-1])], Inches(0.6), Inches(1.6), Inches(6.0), Inches(5.2), legend=False, number_format="#,##0.0")
                rows = [["Advertiser", "GRP share", "CPRP"]] + [[r["name"], f"{r['grp_share_pct']:.1f}%", f"LKR {r['cprp']:,.0f}"] for r in ga]
                self.table(s, rows, Inches(6.9), Inches(1.7), Inches(5.9), size=11)
            if d["alerts"]["items"]:
                s = self.slide("Competitor movements", "Latest week compared with the week before")
                self.bullets(s, [a["text"] for a in d["alerts"]["items"][:7]], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), 15)
        elif rep.kind == "tv":
            top = d["top_programmes"][:10]
            s = self.slide("Top programmes", "Average TVR across airings")
            self.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [f"{r['program']} ({r['channel']})" for r in top][::-1], [("Average TVR", [r["avg_tvr"] for r in top][::-1])], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), legend=False, number_format="0.00")
            s = self.slide("Channel comparison", "Average TVR by channel")
            self.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, [r["channel"] for r in d["by_channel"]], [("Average TVR", [r["avg_tvr"] for r in d["by_channel"]])], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), legend=False, number_format="0.00")
        else:
            tr = d["trend"]
            from .export_pdf import primary_series
            main = primary_series(d)
            s = self.slide(f"{main['label']}: performance trend", f"{tr['granularity'].title()} values")
            self.chart(s, XL_CHART_TYPE.LINE, tr["dates"], [(main["label"], main["values"])], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), legend=False)
            if d.get("segments"):
                first = next((m for m in d["metrics"] if m["agg"] == "sum"), None)
                if first:
                    segs = d["segments"][:10]
                    s = self.slide(f"{first['label']} by segment", "Share of the total")
                    self.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [r["segment"] for r in segs][::-1], [(first["label"], [r[first["key"]] or 0 for r in segs][::-1])], Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), legend=False)


def _p(v: float | None, signed: bool = False) -> str:
    if v is None:
        return "n/a"
    return f"{v:+.0f}%" if signed else f"{v:.0f}%"


def _fmt(v: Any, f: str) -> str:
    if v is None:
        return "n/a"
    if isinstance(v, str):
        return v
    if f == "lkr":
        return f"LKR {v / 1e6:,.1f}M" if abs(v) >= 1e6 else f"LKR {v:,.0f}"
    if f == "pct":
        return f"{v * 100:.1f}%"
    if f in ("dec1", "dec2"):
        return f"{v:,.2f}"
    return f"{v:,.0f}"


def build_pptx(rep: Report) -> bytes:
    return Deck(rep).build()
